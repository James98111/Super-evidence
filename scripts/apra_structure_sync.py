#!/usr/bin/env python3
"""Build a compact, consumer-facing APRA super market structure snapshot.

The APRA Product Structure workbook is pathway-level. This script deliberately
collapses those rows into unique investment options, while retaining the
products/menus through which each option is available.

Outputs:
  data/apra-market/funds-index.json
  data/apra-market/funds/<slug>.json
  data/apra-market/structure-schema.json
  data/apra-market/structure-manifest.json
"""
from __future__ import annotations

import hashlib
import json
import re
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "apra-market"
FUNDS = OUT / "funds"
URL = "https://www.apra.gov.au/system/files/2026-06/Quarterly%20Superannution%20Product%20Statistics%20-%20Product%20Structure.xlsx"
TMP = Path("/tmp/apra-product-structure.xlsx")
UA = "SuperEvidence/0.2 (+public Australian super research)"
REPORTING_PERIOD = "2026-03-31"

DIRECT_CATEGORIES = {
    "direct cash account",
    "direct term deposit",
    "direct fixed income instrument",
    "direct shares",
    "direct hybrid security",
}


def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").replace("\n", " ")).strip().lower()


def clean(v) -> str:
    return "" if v is None else str(v).strip()


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower().replace("&", " and ")).strip("-") or "fund"


def fetch():
    req = urllib.request.Request(URL, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=120) as r, TMP.open("wb") as f:
        while True:
            b = r.read(1024 * 1024)
            if not b:
                break
            f.write(b)
    print("downloaded", TMP.stat().st_size, "bytes", flush=True)


def score(row) -> int:
    ns = [norm(x) for x in row]
    terms = ["rse", "superannuation product", "investment menu", "investment option", "identifier"]
    return sum(any(t in x for t in terms) for x in ns)


def find(headers, *terms):
    mapping = {h: norm(h) for h in headers}
    for term in terms:
        t = norm(term)
        for h, n in mapping.items():
            if n == t:
                return h
    for term in terms:
        t = norm(term)
        candidates = [h for h, n in mapping.items() if t in n]
        if candidates:
            return min(candidates, key=lambda h: len(mapping[h]))
    return None


def rowdict(headers, values):
    vals = list(values) + [None] * max(0, len(headers) - len(values))
    return {headers[i]: vals[i] for i in range(len(headers))}


def classify(category: str) -> str:
    c = norm(category)
    return "direct_asset" if c in DIRECT_CATEGORIES or c.startswith("direct ") else "managed_option"


def indicator(v):
    s = norm(v)
    if s in {"yes", "y", "true", "1"}:
        return True
    if s in {"no", "n", "false", "0"}:
        return False
    return None


def number(v):
    if v is None or clean(v) in {"", "-"}:
        return None
    try:
        return float(str(v).replace(",", "").replace("$", ""))
    except ValueError:
        return None


def headers_for(ws):
    it = ws.iter_rows(values_only=True)
    probe = []
    for _ in range(40):
        try:
            probe.append(next(it))
        except StopIteration:
            break
    if not probe:
        return None, None, None
    idx = max(range(len(probe)), key=lambda i: score(probe[i]))
    if score(probe[idx]) < 2:
        return None, None, None
    headers = [clean(x) or f"column_{i+1}" for i, x in enumerate(probe[idx])]
    return headers, probe[idx + 1 :], it


def parse_product_sheet(ws, schemas):
    headers, buffered, it = headers_for(ws)
    if not headers:
        schemas.append({"sheet": ws.title, "skipped": True, "reason": "no product header detected"})
        return []
    cols = {
        "fund": find(headers, "RSE name"),
        "rse_abn": find(headers, "RSE ABN"),
        "product_id": find(headers, "Product Identifier", "Superannuation Product Identifier"),
        "product": find(headers, "Superannuation Product Name"),
        "product_type": find(headers, "Superannuation Product Type"),
        "product_category": find(headers, "Superannuation Product Category Type"),
        "phase": find(headers, "Superannuation Product Phase Type"),
        "open_employers": find(headers, "Open To New Employers Superannuation Product Indicator"),
        "open_members": find(headers, "Open To New Members Superannuation Product Indicator"),
        "open_public": find(headers, "Open To Public Superannuation Product Indicator"),
        "inception": find(headers, "Superannuation Product Inception Date"),
        "end": find(headers, "Superannuation Product End Date"),
        "member_accounts": find(headers, "Member Accounts (rounded to nearest 10)"),
        "member_assets": find(headers, "Member Assets"),
        "pds_url": find(headers, "Product Disclosure Statement URL"),
    }
    schemas.append({"sheet": ws.title, "headers": headers, "mapped": cols})
    out = []
    for vals in list(buffered) + list(it):
        d = rowdict(headers, vals)
        fund = clean(d.get(cols["fund"])) if cols["fund"] else ""
        product = clean(d.get(cols["product"])) if cols["product"] else ""
        product_id = clean(d.get(cols["product_id"])) if cols["product_id"] else ""
        if not fund or (not product and not product_id):
            continue
        out.append({
            "fund": fund,
            "rse_abn": clean(d.get(cols["rse_abn"])) if cols["rse_abn"] else "",
            "product_id": product_id,
            "product_name": product,
            "product_type": clean(d.get(cols["product_type"])) if cols["product_type"] else "",
            "product_category": clean(d.get(cols["product_category"])) if cols["product_category"] else "",
            "phase": clean(d.get(cols["phase"])) if cols["phase"] else "",
            "open_to_new_employers": indicator(d.get(cols["open_employers"])) if cols["open_employers"] else None,
            "open_to_new_members": indicator(d.get(cols["open_members"])) if cols["open_members"] else None,
            "open_to_public": indicator(d.get(cols["open_public"])) if cols["open_public"] else None,
            "inception_date": clean(d.get(cols["inception"])) if cols["inception"] else "",
            "end_date": clean(d.get(cols["end"])) if cols["end"] else "",
            "member_accounts": number(d.get(cols["member_accounts"])) if cols["member_accounts"] else None,
            "member_assets": number(d.get(cols["member_assets"])) if cols["member_assets"] else None,
            "pds_url": clean(d.get(cols["pds_url"])) if cols["pds_url"] else "",
        })
    print(ws.title, len(out), "product rows", flush=True)
    return out


def parse_pathway_sheet(ws, schemas):
    headers, buffered, it = headers_for(ws)
    if not headers:
        schemas.append({"sheet": ws.title, "skipped": True, "reason": "no pathway header detected"})
        return []
    cols = {
        "pathway_id": find(headers, "Pathway Identifier"),
        "fund": find(headers, "RSE name"),
        "rse_abn": find(headers, "RSE ABN"),
        "product_id": find(headers, "Product Identifier", "Superannuation Product Identifier"),
        "product": find(headers, "Superannuation Product Name"),
        "product_type": find(headers, "Superannuation Product Type"),
        "product_category": find(headers, "Superannuation Product Category Type"),
        "phase": find(headers, "Superannuation Product Phase Type"),
        "menu_id": find(headers, "Investment Menu Identifier"),
        "menu": find(headers, "Investment Menu Name"),
        "menu_type": find(headers, "Investment Menu Type"),
        "menu_open": find(headers, "Investment menu open to new members"),
        "option_id": find(headers, "Investment Option Identifier"),
        "option": find(headers, "Investment Option / Lifecycle Stage Name", "Investment Option Name"),
        "option_type": find(headers, "Investment option type", "Investment Option Type"),
        "category": find(headers, "Investment Option Category Type"),
        "apir": find(headers, "Investment Option APIR Code"),
        "exchange": find(headers, "Exchange Code"),
        "ticker": find(headers, "Ticker Symbol"),
        "open_option": find(headers, "Open To New Members Investment Option Indicator"),
        "inception": find(headers, "Investment Option Inception Date"),
        "end": find(headers, "Investment Option End Date"),
        "pathway_accounts": find(headers, "Investment Pathway Member Accounts (rounded to nearest 10)"),
        "pathway_assets": find(headers, "Investment Pathway Member Assets"),
        "option_accounts": find(headers, "Investment Option Member Accounts (All investment pathways)"),
        "option_assets": find(headers, "Investment Option Member Assets (All investment pathways)"),
    }
    schemas.append({"sheet": ws.title, "headers": headers, "mapped": cols})
    out = []
    for vals in list(buffered) + list(it):
        d = rowdict(headers, vals)
        fund = clean(d.get(cols["fund"])) if cols["fund"] else ""
        option = clean(d.get(cols["option"])) if cols["option"] else ""
        option_id = clean(d.get(cols["option_id"])) if cols["option_id"] else ""
        if not fund or (not option and not option_id):
            continue
        out.append({
            "pathway_id": clean(d.get(cols["pathway_id"])) if cols["pathway_id"] else "",
            "fund": fund,
            "rse_abn": clean(d.get(cols["rse_abn"])) if cols["rse_abn"] else "",
            "product_id": clean(d.get(cols["product_id"])) if cols["product_id"] else "",
            "product_name": clean(d.get(cols["product"])) if cols["product"] else "",
            "product_type": clean(d.get(cols["product_type"])) if cols["product_type"] else "",
            "product_category": clean(d.get(cols["product_category"])) if cols["product_category"] else "",
            "phase": clean(d.get(cols["phase"])) if cols["phase"] else "",
            "menu_id": clean(d.get(cols["menu_id"])) if cols["menu_id"] else "",
            "menu_name": clean(d.get(cols["menu"])) if cols["menu"] else "",
            "menu_type": clean(d.get(cols["menu_type"])) if cols["menu_type"] else "",
            "menu_open_to_new_members": indicator(d.get(cols["menu_open"])) if cols["menu_open"] else None,
            "option_id": option_id,
            "option_name": option,
            "option_type": clean(d.get(cols["option_type"])) if cols["option_type"] else "",
            "option_category": clean(d.get(cols["category"])) if cols["category"] else "",
            "apir_code": clean(d.get(cols["apir"])) if cols["apir"] else "",
            "exchange_code": clean(d.get(cols["exchange"])) if cols["exchange"] else "",
            "ticker_symbol": clean(d.get(cols["ticker"])) if cols["ticker"] else "",
            "open_to_new_members": indicator(d.get(cols["open_option"])) if cols["open_option"] else None,
            "inception_date": clean(d.get(cols["inception"])) if cols["inception"] else "",
            "end_date": clean(d.get(cols["end"])) if cols["end"] else "",
            "pathway_member_accounts": number(d.get(cols["pathway_accounts"])) if cols["pathway_accounts"] else None,
            "pathway_member_assets": number(d.get(cols["pathway_assets"])) if cols["pathway_assets"] else None,
            "option_member_accounts": number(d.get(cols["option_accounts"])) if cols["option_accounts"] else None,
            "option_member_assets": number(d.get(cols["option_assets"])) if cols["option_assets"] else None,
        })
    print(ws.title, len(out), "pathway rows", flush=True)
    return out


def compact_options(rows):
    grouped = {}
    for r in rows:
        # Investment Option Identifier is the regulator's option identity. Use the
        # name only as a defensive fallback when an identifier is absent.
        key = (r["fund"], r["option_id"] or f"name::{r['option_name']}")
        if key not in grouped:
            grouped[key] = {
                "fund": r["fund"],
                "rse_abn": r["rse_abn"],
                "option_id": r["option_id"],
                "option_name": r["option_name"],
                "aliases": [],
                "option_type": r["option_type"],
                "option_category": r["option_category"],
                "consumer_group": classify(r["option_category"]),
                "apir_code": r["apir_code"],
                "exchange_code": r["exchange_code"],
                "ticker_symbol": r["ticker_symbol"],
                "open_to_new_members": r["open_to_new_members"],
                "inception_date": r["inception_date"],
                "end_date": r["end_date"],
                "member_accounts": r["option_member_accounts"],
                "member_assets": r["option_member_assets"],
                "pathways": [],
            }
        o = grouped[key]
        if r["option_name"] and r["option_name"] != o["option_name"] and r["option_name"] not in o["aliases"]:
            o["aliases"].append(r["option_name"])
        # Prefer non-empty metadata; never average or manufacture it.
        for k in ["option_type", "option_category", "apir_code", "exchange_code", "ticker_symbol", "inception_date", "end_date"]:
            if not o.get(k) and r.get(k):
                o[k] = r[k]
        if o.get("open_to_new_members") is None and r.get("open_to_new_members") is not None:
            o["open_to_new_members"] = r["open_to_new_members"]
        for k in ["member_accounts", "member_assets"]:
            if o.get(k) is None and r.get("option_" + k) is not None:
                o[k] = r["option_" + k]
        p = {
            "pathway_id": r["pathway_id"],
            "product_id": r["product_id"],
            "product_name": r["product_name"],
            "product_type": r["product_type"],
            "product_category": r["product_category"],
            "phase": r["phase"],
            "menu_id": r["menu_id"],
            "menu_name": r["menu_name"],
            "menu_type": r["menu_type"],
            "menu_open_to_new_members": r["menu_open_to_new_members"],
            "pathway_member_accounts": r["pathway_member_accounts"],
            "pathway_member_assets": r["pathway_member_assets"],
        }
        # Pathway ID is unique where present. If absent, avoid duplicate product/menu tuples.
        pk = p["pathway_id"] or (p["product_id"], p["menu_id"], p["phase"])
        existing = {x["pathway_id"] or (x["product_id"], x["menu_id"], x["phase"]) for x in o["pathways"]}
        if pk not in existing:
            o["pathways"].append(p)
    return list(grouped.values())


def main():
    fetch()
    digest = hashlib.sha256(TMP.read_bytes()).hexdigest()
    wb = load_workbook(TMP, read_only=True, data_only=True)
    schemas = []
    products = []
    pathways = []
    for ws in wb.worksheets:
        if ws.title == "Table 1a":
            products.extend(parse_product_sheet(ws, schemas))
        elif ws.title == "Table 1b":
            pathways.extend(parse_pathway_sheet(ws, schemas))
        else:
            headers, _, _ = headers_for(ws)
            if headers:
                schemas.append({"sheet": ws.title, "headers": headers, "mapped": {}})
            else:
                schemas.append({"sheet": ws.title, "skipped": True, "reason": "no structure header detected"})

    options = compact_options(pathways)
    products_by_fund = defaultdict(list)
    options_by_fund = defaultdict(list)
    for p in products:
        products_by_fund[p["fund"]].append(p)
    for o in options:
        options_by_fund[o["fund"]].append(o)

    # Funds that appear only in one table are still retained.
    fund_names = sorted(set(products_by_fund) | set(options_by_fund), key=str.casefold)
    OUT.mkdir(parents=True, exist_ok=True)
    FUNDS.mkdir(parents=True, exist_ok=True)

    index = []
    used = {}
    for name in fund_names:
        ps = products_by_fund[name]
        os = options_by_fund[name]
        base = slug(name)
        s = base
        if s in used and used[s] != name:
            s = f"{base}-{hashlib.sha1(name.encode()).hexdigest()[:7]}"
        used[s] = name
        managed = [o for o in os if o["consumer_group"] == "managed_option"]
        direct = [o for o in os if o["consumer_group"] == "direct_asset"]
        pathway_count = sum(len(o["pathways"]) for o in os)
        open_public_products = [p for p in ps if p["open_to_public"] is True and p["open_to_new_members"] is not False and not p["end_date"]]
        rse_abn = next((p["rse_abn"] for p in ps if p["rse_abn"]), next((o["rse_abn"] for o in os if o["rse_abn"]), None))
        payload = {
            "schema_version": 2,
            "source": {
                "organisation": "APRA",
                "publication": "Quarterly Superannuation Product Statistics — Product Structure",
                "reporting_period": REPORTING_PERIOD,
                "url": URL,
            },
            "fund": {
                "name": name,
                "rse_abn": rse_abn,
                "slug": s,
                "open_public_product_count": len(open_public_products),
            },
            "summary": {
                "products": len(ps),
                "open_public_products": len(open_public_products),
                "unique_managed_options": len(managed),
                "unique_direct_assets": len(direct),
                "unique_options_total": len(os),
                "pathways": pathway_count,
            },
            "products": sorted(ps, key=lambda p: (p["phase"], p["product_name"])),
            "options": sorted(os, key=lambda o: (o["consumer_group"], o["option_category"], o["option_name"])),
        }
        (FUNDS / f"{s}.json").write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        index.append({"name": name, "slug": s, "rse_abn": rse_abn, **payload["summary"]})

    if len(index) < 20:
        raise RuntimeError(f"Only {len(index)} funds parsed; refusing publication")
    unique_total = sum(x["unique_options_total"] for x in index)
    if unique_total < 1000:
        raise RuntimeError(f"Only {unique_total} unique options parsed; refusing publication")
    generated = datetime.now(timezone.utc).isoformat()
    market_index = {
        "schema_version": 2,
        "source": "APRA QSPS Product Structure",
        "reporting_period": REPORTING_PERIOD,
        "generated_at": generated,
        "fund_count": len(index),
        "unique_managed_options": sum(x["unique_managed_options"] for x in index),
        "unique_direct_assets": sum(x["unique_direct_assets"] for x in index),
        "unique_options_total": unique_total,
        "pathways": sum(x["pathways"] for x in index),
        "funds": index,
    }
    (OUT / "funds-index.json").write_text(json.dumps(market_index, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    (OUT / "structure-schema.json").write_text(json.dumps(schemas, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "structure-manifest.json").write_text(json.dumps({
        "url": URL,
        "sha256": digest,
        "bytes": TMP.stat().st_size,
        "generated_at": generated,
        "fund_count": len(index),
        "unique_options_total": unique_total,
        "pathways": market_index["pathways"],
        "parser_version": "apra-structure-v2",
    }, indent=2), encoding="utf-8")
    print(
        "VALIDATED",
        len(index), "funds",
        market_index["unique_managed_options"], "managed options",
        market_index["unique_direct_assets"], "direct assets",
        market_index["pathways"], "pathways",
        flush=True,
    )


if __name__ == "__main__":
    main()
