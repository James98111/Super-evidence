#!/usr/bin/env python3
from __future__ import annotations
import json, urllib.request
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data'/'apra-market'/'performance-real-samples.json'
URL='https://www.apra.gov.au/system/files/2026-06/Quarterly%20Superannuation%20Product%20Publication%20-%20Performance_0.xlsx'
TMP=Path('/tmp/apra-performance.xlsx')
SHEETS=['Table 4a','Table 5a','Table 6a','Table 7a','Table 8a','Table 8b','Table 8c','Table 8d','Table 9','Table 9a']
UA='SuperEvidence/0.3 (+public Australian super research)'

def clean(v):
    if v is None:return None
    if hasattr(v,'isoformat'):return v.isoformat()
    return v

def fetch():
    req=urllib.request.Request(URL,headers={'User-Agent':UA})
    with urllib.request.urlopen(req,timeout=180) as r, TMP.open('wb') as f:
        while True:
            b=r.read(1024*1024)
            if not b:break
            f.write(b)

def header_row(ws):
    for i,row in enumerate(ws.iter_rows(min_row=1,max_row=12,values_only=True),1):
        vals=[str(v or '') for v in row]
        if any('Investment Option Identifier' in v or 'Pathway Identifier' in v for v in vals):return i
    raise RuntimeError(f'header not found {ws.title}')

def main():
    fetch();wb=load_workbook(TMP,read_only=True,data_only=True)
    result={'source':URL,'sheets':{}}
    for name in SHEETS:
        ws=wb[name];hr=header_row(ws)
        headers=[str(v).strip() if v is not None and str(v).strip() else f'column_{i+1}' for i,v in enumerate(next(ws.iter_rows(min_row=hr,max_row=hr,values_only=True)))]
        samples=[]
        for row in ws.iter_rows(min_row=hr+1,values_only=True):
            d={headers[i]:clean(row[i]) if i<len(row) else None for i in range(len(headers))}
            ident=d.get('Pathway Identifier') or d.get('Investment Option Identifier')
            period=d.get('Period')
            if not ident:continue
            if not period or str(period).startswith('*'):continue
            samples.append(d)
            if len(samples)>=3:break
        result['sheets'][name]={'header_row':hr,'samples':samples}
        print(name,len(samples),flush=True)
    OUT.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')

if __name__=='__main__':main()
