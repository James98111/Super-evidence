#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, re, urllib.request
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data'/'apra-market'
URL='https://www.apra.gov.au/system/files/2026-06/Quarterly%20Superannuation%20Product%20Publication%20-%20Performance_0.xlsx'
TMP=Path('/tmp/apra-performance.xlsx')
UA='SuperEvidence/0.2 (+public Australian super research)'

def clean(v): return '' if v is None else str(v).strip()
def norm(v): return re.sub(r'\s+',' ',clean(v).replace('\n',' ')).strip().lower()
def score(row):
    ns=[norm(x) for x in row]
    terms=['rse','investment option','pathway identifier','return','performance','fee','asset allocation','strategic']
    return sum(any(t in x for t in terms) for x in ns)

def main():
    req=urllib.request.Request(URL,headers={'User-Agent':UA})
    with urllib.request.urlopen(req,timeout=120) as r, TMP.open('wb') as f:
        while True:
            b=r.read(1024*1024)
            if not b: break
            f.write(b)
    print('downloaded',TMP.stat().st_size,'bytes',flush=True)
    wb=load_workbook(TMP,read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        it=ws.iter_rows(values_only=True)
        probe=[]
        for _ in range(60):
            try: probe.append(next(it))
            except StopIteration: break
        if not probe:
            continue
        idx=max(range(len(probe)),key=lambda i:score(probe[i]))
        headers=[clean(x) or f'column_{i+1}' for i,x in enumerate(probe[idx])]
        header_score=score(probe[idx])
        samples=[]
        if header_score>=2:
            candidates=list(probe[idx+1:])
            for _ in range(8):
                try:candidates.append(next(it))
                except StopIteration:break
            for vals in candidates:
                if any(clean(v) for v in vals):
                    row={headers[i]:clean(vals[i]) if i<len(vals) else '' for i in range(len(headers))}
                    if any(row.values()): samples.append(row)
                    if len(samples)>=3:break
        sheets.append({'sheet':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'header_score':header_score,'header_row_1_based':idx+1,'headers':headers,'samples':samples})
        print(ws.title,'rows',ws.max_row,'cols',ws.max_column,'score',header_score,flush=True)
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'performance-schema.json').write_text(json.dumps({'url':URL,'sha256':hashlib.sha256(TMP.read_bytes()).hexdigest(),'bytes':TMP.stat().st_size,'sheets':sheets},ensure_ascii=False,indent=2),encoding='utf-8')

if __name__=='__main__':main()
