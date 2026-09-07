#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MARKET = ROOT / 'data' / 'apra-market'
METRICS_DIR = MARKET / 'metrics'
STRUCTURE_INDEX = MARKET / 'funds-index.json'
URL = 'https://www.apra.gov.au/system/files/2026-06/Quarterly%20Superannuation%20Product%20Publication%20-%20Performance_0.xlsx'
TMP = Path('/tmp/apra-performance.xlsx')
UA = 'SuperEvidence/0.4 (+public Australian super research)'

PERF_SHEETS = {
    'Table 4a': ('mysuper', 50000, 'accumulation'),
    'Table 4b': ('mysuper', 100000, 'accumulation'),
    'Table 4c': ('mysuper', 250000, 'accumulation'),
    'Table 5a': ('choice_non_platform', 50000, 'accumulation'),
    'Table 5b': ('choice_non_platform', 100000, 'accumulation'),
    'Table 5c': ('choice_non_platform', 250000, 'accumulation'),
    'Table 6a': ('choice_platform', 50000, 'accumulation'),
    'Table 6b': ('choice_platform', 100000, 'accumulation'),
    'Table 6c': ('choice_platform', 250000, 'accumulation'),
    'Table 7a': ('retirement', 50000, 'retirement'),
    'Table 7b': ('retirement', 100000, 'retirement'),
    'Table 7c': ('retirement', 250000, 'retirement'),
    'Table 7d': ('retirement', 500000, 'retirement'),
}

STRATEGY_SHEETS = {
    'Table 8a': ('mysuper', 'accumulation'),
    'Table 8b': ('choice_non_platform', 'accumulation'),
    'Table 8c': ('choice_platform', 'accumulation'),
    'Table 8d': ('retirement', 'retirement'),
}


def clean(v):
    if v is None:
        return None
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    if isinstance(v, str):
        v = re.sub(r'\s+', ' ', v).strip()
        return v or None
    return v


def norm(v):
    return re.sub(r'\s+', ' ', str(v or '').replace('\n', ' ')).strip().lower()


def pct_value(v):
    if v is None or v == '':
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    # APRA percentage fields are stored as decimal fractions in the workbook.
    return round(x * 100.0, 6)


def number(v):
    if v is None or v == '':
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def fetch():
    req = urllib.request.Request(URL, headers={'User-Agent': UA})
    with urllib.request.urlopen(req, timeout=180) as r, TMP.open('wb') as f:
        while True:
            b = r.read(1024 * 1024)
            if not b:
                break
            f.write(b)
    print('downloaded', TMP.stat().st_size, 'bytes', flush=True)


def header_row(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
        vals = [str(v or '') for v in row]
        if any('Investment Option Identifier' in v or 'Pathway Identifier' in v for v in vals):
            return i
    raise RuntimeError(f'No header row found for {ws.title}')


def rows(ws):
    hr = header_row(ws)
    header_vals = next(ws.iter_rows(min_row=hr, max_row=hr, values_only=True))
    headers = [str(v).strip() if v is not None and str(v).strip() else f'column_{i+1}' for i, v in enumerate(header_vals)]
    for raw in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {headers[i]: clean(raw[i]) if i < len(raw) else None for i in range(len(headers))}
        yield d


def pick(d, *needles):
    needles = [norm(x) for x in needles]
    for key, value in d.items():
        nk = norm(key)
        if all(x in nk for x in needles):
            return value
    return None


def return_value(d, years, member=False):
    y = f'{years}-year'
    # APRA labels vary by MySuper / Choice / platform / retirement. Match the
    # horizon plus annualised, then distinguish member net return from the
    # investment-return metric.
    candidates = []
    for k, v in d.items():
        nk = norm(k)
        if y not in nk or 'annualised' not in nk:
            continue
        if member:
            if 'net return' in nk and 'investment return' not in nk:
                candidates.append(v)
        else:
            if ('investment return' in nk or 'gross investment return net of fees' in nk) and 'net return gross of tax' not in nk:
                candidates.append(v)
    for v in candidates:
        x = pct_value(v)
        if x is not None:
            return x
    return None


def performance_basis(sheet):
    if sheet.startswith('Table 6'):
        return 'gross_investment_return_net_of_fees'
    if sheet.startswith('Table 7'):
        return 'investment_return_net_of_fees'
    return 'net_investment_return'


def date_key(v):
    return str(v or '')


def parse_performance(wb):
    latest = {}
    for sheet, (segment, representative_balance, phase) in PERF_SHEETS.items():
        ws = wb[sheet]
        seen = 0
        for d in rows(ws):
            pathway_id = clean(d.get('Pathway Identifier'))
            option_id = clean(d.get('Investment Option Identifier'))
            fund = clean(d.get('RSE name'))
            period = clean(d.get('Period'))
            if not pathway_id or not option_id or not fund or not period:
                continue
            seen += 1
            key = (sheet, pathway_id)
            if key in latest and date_key(latest[key]['reporting_period']) >= date_key(period):
                continue
            rec = {
                'pathway_id': pathway_id,
                'option_id': option_id,
                'fund': fund,
                'product_id': clean(d.get('Product Identifier')),
                'product_name': clean(d.get('Superannuation Product Name')),
                'menu_id': clean(d.get('Investment Menu Identifier')),
                'menu_name': clean(d.get('Investment Menu Name')),
                'option_name': clean(d.get('Investment Option / Lifecycle Stage Name')),
                'option_type': clean(d.get('Investment Option Type')),
                'option_category': clean(d.get('Investment Option Category')),
                'segment': segment,
                'phase': phase,
                'representative_balance': representative_balance,
                'reporting_period': period,
                'investment_return_basis': performance_basis(sheet),
                'investment_return_1y_pct': return_value(d, 1, member=False),
                'investment_return_3y_pct': return_value(d, 3, member=False),
                'investment_return_5y_pct': return_value(d, 5, member=False),
                'investment_return_10y_pct': return_value(d, 10, member=False),
                'member_net_return_1y_pct': return_value(d, 1, member=True),
                'member_net_return_3y_pct': return_value(d, 3, member=True),
                'member_net_return_5y_pct': return_value(d, 5, member=True),
                'member_net_return_10y_pct': return_value(d, 10, member=True),
                'volatility_10y_pct': pct_value(pick(d, 'volatility', '10 year')),
                'investment_fees_costs_pct': pct_value(pick(d, 'total investment fees and costs', 'rep member')),
                'transaction_fees_costs_pct': pct_value(pick(d, 'total transaction fees and costs', 'rep member')),
                'administration_fees_costs_pct': pct_value(pick(d, 'total administration fees and costs', 'rep member')),
                'total_fees_costs_pct': pct_value(pick(d, 'total fees and costs', 'rep member')),
                'total_fees_costs_taxes_pct': pct_value(pick(d, 'total fees, costs and taxes', 'rep member')),
            }
            latest[key] = rec
        print(sheet, seen, 'performance rows scanned', flush=True)
    return list(latest.values())


def parse_strategy(wb):
    latest = {}
    for sheet, (segment, phase) in STRATEGY_SHEETS.items():
        ws = wb[sheet]
        seen = 0
        for d in rows(ws):
            option_id = clean(d.get('Investment Option Identifier'))
            fund = clean(d.get('RSE name'))
            period = clean(d.get('Period'))
            if not option_id or not fund or not period:
                continue
            seen += 1
            key = (sheet, option_id)
            if key in latest and date_key(latest[key]['reporting_period']) >= date_key(period):
                continue

            allocations = []
            for col, value in d.items():
                ncol = norm(col)
                if 'benchmark asset allocation' not in ncol:
                    continue
                label = re.sub(r'\s*-\s*Benchmark asset allocation\s*$', '', str(col), flags=re.I).strip()
                benchmark = pct_value(value)
                lower = None
                upper = None
                for c2, v2 in d.items():
                    nc2 = norm(c2)
                    if norm(label) not in nc2:
                        continue
                    if 'lower end of asset allocation range' in nc2:
                        lower = pct_value(v2)
                    elif 'upper end of asset allocation range' in nc2:
                        upper = pct_value(v2)
                if benchmark is not None or lower is not None or upper is not None:
                    allocations.append({'asset_class': label, 'benchmark_pct': benchmark, 'lower_pct': lower, 'upper_pct': upper})

            rec = {
                'option_id': option_id,
                'fund': fund,
                'option_name': clean(d.get('Investment Option / Lifecycle Stage Name') or d.get('Investment Option Name')),
                'option_type': clean(d.get('Investment Option Type')),
                'option_category': clean(d.get('Investment Option Category')),
                'segment': segment,
                'phase': phase,
                'reporting_period': period,
                'growth_weight_pct': pct_value(pick(d, 'growth asset weighting')),
                'growth_band': clean(pick(d, 'growth asset band')),
                'currency_exposure_pct': pct_value(pick(d, 'currency exposure')),
                'investment_horizon_years': number(pick(d, 'investment horizon years number')),
                'return_margin_pct': pct_value(pick(d, 'return margin percent')),
                'return_objective_benchmark': clean(pick(d, 'return objective benchmark text')),
                'risk_label': clean(pick(d, 'level of investment risk label')),
                'negative_return_expectation_20yr': clean(pick(d, 'level of investment risk', '20 year')),
                'allocations': allocations,
            }
            latest[key] = rec
        print(sheet, seen, 'strategy rows scanned', flush=True)
    return list(latest.values())


def main():
    if not STRUCTURE_INDEX.exists():
        raise RuntimeError('Run apra_structure_sync.py first')
    index = json.loads(STRUCTURE_INDEX.read_text(encoding='utf-8'))
    slug_by_fund = {x['name']: x['slug'] for x in index['funds']}

    fetch()
    wb = load_workbook(TMP, read_only=True, data_only=True)
    performance = parse_performance(wb)
    strategy = parse_strategy(wb)

    by_fund = defaultdict(lambda: {'performance': [], 'strategy': []})
    for r in performance:
        by_fund[r['fund']]['performance'].append(r)
    for r in strategy:
        by_fund[r['fund']]['strategy'].append(r)

    METRICS_DIR.mkdir(parents=True, exist_ok=True)
    generated = datetime.now(timezone.utc).isoformat()
    all_growth = []
    strategy_ids = set()
    perf_pathways = set()
    perf_options = set()

    for fund, payload in by_fund.items():
        slug = slug_by_fund.get(fund)
        if not slug:
            continue
        for r in payload['strategy']:
            if r['growth_weight_pct'] is not None:
                all_growth.append(r['growth_weight_pct'])
            strategy_ids.add(r['option_id'])
        for r in payload['performance']:
            perf_pathways.add(r['pathway_id'])
            perf_options.add(r['option_id'])
        out = {
            'schema_version': 1,
            'source': {'organisation': 'APRA', 'publication': 'Quarterly Superannuation Product Statistics — Performance', 'url': URL},
            'generated_at': generated,
            'fund': fund,
            'performance': sorted(payload['performance'], key=lambda x: (x['option_name'] or '', x['representative_balance'], x['segment'])),
            'strategy': sorted(payload['strategy'], key=lambda x: (x['option_name'] or '', x['segment'])),
        }
        (METRICS_DIR / f'{slug}.json').write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')

    public_funds = sum(1 for x in index['funds'] if x.get('open_public_products', 0) > 0)
    stats = {
        'schema_version': 1,
        'generated_at': generated,
        'structure_reporting_period': index.get('reporting_period'),
        'fund_entities': index.get('fund_count'),
        'fund_entities_with_public_products': public_funds,
        'unique_managed_options': index.get('unique_managed_options'),
        'unique_direct_assets': index.get('unique_direct_assets'),
        'unique_option_records': index.get('unique_options_total'),
        'investment_pathways': index.get('pathways'),
        'options_with_strategy_metrics': len(strategy_ids),
        'options_with_performance_metrics': len(perf_options),
        'performance_pathways': len(perf_pathways),
        'growth_exposure_min_pct': round(min(all_growth), 2) if all_growth else None,
        'growth_exposure_max_pct': round(max(all_growth), 2) if all_growth else None,
        'source_note': 'Market structure is APRA QSPS March 2026. Performance and strategy coverage is drawn from APRA quarterly performance tables and varies by product type and reporting history.'
    }
    (MARKET / 'market-stats.json').write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding='utf-8')

    if stats['fund_entities'] < 50 or stats['unique_option_records'] < 10000:
        raise RuntimeError('Structure coverage validation failed')
    if len(strategy_ids) < 500 or len(perf_options) < 500:
        raise RuntimeError(f'Metric coverage unexpectedly low: strategy={len(strategy_ids)}, performance={len(perf_options)}')
    print('VALIDATED', json.dumps(stats, sort_keys=True), flush=True)

if __name__ == '__main__':
    main()
